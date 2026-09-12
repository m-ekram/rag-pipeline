"""Build Dissertation.docx from the Markdown chapters in ./chapters.

    ragenv311\\Scripts\\python docs/dissertation/build_docx.py

The chapters stay the source of truth (reviewable, diffable); this script only
typesets them. Supported Markdown: `#`/`##`/`###` headings, paragraphs,
`-` and `1.` lists, pipe tables, fenced code, `![caption](path)` figures, and
inline **bold**, *italic* and `code`. Word fills in the table of contents when
the document is opened (it asks to update fields).
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
CHAPTERS = HERE / "chapters"
OUTPUT = HERE / "Dissertation.docx"

META = {
    "title": "Sanchay: Optimising a Local Retrieval-Augmented Generation System "
             "for Heterogeneous Government Documents",
    "subtitle": "Hybrid retrieval, OCR and grounded generation over the Patna Master Plan, "
                "Bihar electoral rolls and research papers",
    "author": "[Author name]",
    "degree": "[Degree programme]",
    "institution": "[University / Department]",
    "supervisor": "[Supervisor]",
}

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
_TABLE_RULE = re.compile(r"^\|?\s*:?-{3,}")


def _add_runs(paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _field(paragraph, instruction: str) -> None:
    """Insert a Word field (TOC, PAGE) that Word evaluates when opened."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field." if "TOC" in instruction else "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, placeholder, end):
        run._r.append(element)


def _styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(6)
    for level, size in ((1, 18), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    # Ask Word to refresh fields (the TOC) when the file is opened.
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _title_page(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(META["title"])
    run.bold = True
    run.font.size = Pt(20)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(META["subtitle"]).italic = True
    for _ in range(4):
        doc.add_paragraph()
    for line in (META["author"], META["degree"], META["institution"],
                 f"Supervisor: {META['supervisor']}", date.today().strftime("%B %Y")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = doc.add_paragraph()
    heading.add_run("Table of Contents").bold = True
    _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _footer_page_numbers(doc: Document) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(p, "PAGE")


def _table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows
             if not _TABLE_RULE.match(r.strip())]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing = 1.0
            _add_runs(para, row[j] if j < len(row) else "")
            for run in para.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def _figure(doc: Document, caption: str, path: str, figure_no: int) -> None:
    image = (CHAPTERS / path).resolve()
    if not image.exists():
        doc.add_paragraph(f"[Missing figure: {path}]")
        return
    doc.add_picture(str(image), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Figure {figure_no}. {caption}")
    run.italic = True
    run.font.size = Pt(10)


def render(doc: Document, text: str, state: dict) -> None:
    lines = text.splitlines()
    i = 0
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, " ".join(s.strip() for s in paragraph))
            paragraph.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            if level == 1 and state["chapters"]:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if level == 1:
                state["chapters"] += 1
            doc.add_heading(title, level=min(level, 3))
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows)
            continue
        elif re.match(r"^!\[(.*)\]\((.*)\)$", stripped):
            flush()
            caption, path = re.match(r"^!\[(.*)\]\((.*)\)$", stripped).groups()
            state["figures"] += 1
            _figure(doc, caption, path, state["figures"])
        elif re.match(r"^(-|\*)\s+", stripped):
            flush()
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^(-|\*)\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            flush()
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", stripped))
        elif not stripped:
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()


def main() -> int:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)
    _styles(doc)
    _title_page(doc)
    state = {"chapters": 0, "figures": 0}
    chapters = sorted(CHAPTERS.glob("*.md"))
    if not chapters:
        print(f"No chapters in {CHAPTERS}", file=sys.stderr)
        return 1
    for chapter in chapters:
        render(doc, chapter.read_text(encoding="utf-8"), state)
    _footer_page_numbers(doc)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(chapters)} chapters, {state['figures']} figures)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
