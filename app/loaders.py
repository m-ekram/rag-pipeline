"""Turn a directory of mixed documents into LangChain Documents.

Each returned Document carries enough metadata to build a real citation:
  source    relative path, e.g. "manuals/networking.pdf"
  page      1-based page number (PDFs only)
  title     human-readable document title
  section   PDF bookmark path in force on that page, when the PDF has one

PDFs come back one Document per page so page numbers survive; the chunker
re-joins them so passages are not severed at page breaks.
"""

from __future__ import annotations

import logging
import os
import re
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

from langchain_core.documents import Document

import config

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".md", ".markdown", ".txt", ".rst", ".html", ".htm", ".docx"}

# Collapse the ragged whitespace PDF extraction leaves behind; it wrecks both
# chunk boundaries and BM25 tokenisation if left in.
_WS = re.compile(r"[ \t ]+")
_BLANKS = re.compile(r"\n{3,}")
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_DIGITS = re.compile(r"\d+")

# Below this much input, spawning worker processes costs more than it saves.
_PARALLEL_MIN_BYTES = 10 * 1024 * 1024


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)  # re-join words split across lines
    text = _WS.sub(" ", text)
    text = _BLANKS.sub("\n\n", text)
    return text.strip()


def _title_from_path(path: Path) -> str:
    return path.stem.replace("_", " ").replace("-", " ").strip()


def _strip_running_lines(pages: list[str]) -> list[str]:
    """Drop page headers and footers that repeat across a PDF.

    Manuals print the chapter title and page number on every page. Left in, a
    10k-page corpus gains thousands of identical fragments that match every
    query about that chapter equally well. A line in the first or last two of a
    page that recurs (digits ignored) on at least 30% of pages is boilerplate.
    """
    if len(pages) < 5:
        return pages

    def edges(lines: list[str]) -> list[int]:
        filled = [i for i, line in enumerate(lines) if line.strip()]
        return filled[:2] + filled[-2:]

    split = [text.split("\n") for text in pages]
    counts: Counter[str] = Counter()
    for lines in split:
        counts.update({_DIGITS.sub("#", lines[i].strip()) for i in edges(lines)})
    threshold = max(3, int(len(pages) * 0.3))
    running = {line for line, n in counts.items() if n >= threshold}
    if not running:
        return pages

    out = []
    for lines in split:
        drop = {i for i in edges(lines) if _DIGITS.sub("#", lines[i].strip()) in running}
        out.append("\n".join(line for i, line in enumerate(lines) if i not in drop).strip())
    return out


def _page_sections(reader, page_count: int) -> dict[int, str]:
    """Map 1-based page -> bookmark path in force there ("Chapter 9 > 9.4 Strings").

    The outline is the only structure a PDF reliably carries; extracted text has
    no heading markup. Depth is capped at three levels, which is where manuals
    stop being navigational and start being noise.
    """
    starts: list[tuple[int, tuple[str, ...]]] = []

    def walk(items, parents: tuple[str, ...]) -> None:
        last = parents
        for item in items:
            if isinstance(item, list):  # children of the preceding entry
                if len(last) < 3:
                    walk(item, last)
                continue
            title = " ".join(str(getattr(item, "title", "") or "").split())
            if not title:
                continue
            last = parents + (title,)
            try:
                starts.append((reader.get_destination_page_number(item) + 1, last))
            except Exception:
                continue

    try:
        walk(reader.outline, ())
    except Exception as exc:  # malformed outlines are common; never fatal
        logger.debug("Outline unreadable: %s", exc)
        return {}

    starts.sort(key=lambda s: s[0])  # stable: a parent stays ahead of its child
    sections: dict[int, str] = {}
    current: tuple[str, ...] | None = None
    j = 0
    for page in range(1, page_count + 1):
        while j < len(starts) and starts[j][0] <= page:
            current = starts[j][1]
            j += 1
        if current:
            sections[page] = " > ".join(current)
    return sections


def _load_pdf(path: Path, source: str) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = (reader.metadata.title if reader.metadata else None) or _title_from_path(path)
    texts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            texts.append(clean_text(page.extract_text() or ""))
        except Exception as exc:  # a single corrupt page shouldn't kill the run
            logger.warning("%s page %d: extraction failed (%s)", source, i, exc)
            texts.append("")

    texts = _strip_running_lines(texts)
    sections = _page_sections(reader, len(texts))

    docs: list[Document] = []
    for i, text in enumerate(texts, start=1):
        if not text:
            continue
        meta = {"source": source, "page": i, "title": str(title)}
        if i in sections:
            meta["section"] = sections[i]
        docs.append(Document(page_content=text, metadata=meta))
    return docs


def _load_html(path: Path, source: str) -> list[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    title = (soup.title.string if soup.title and soup.title.string else None) or _title_from_path(path)
    text = clean_text(soup.get_text("\n"))
    return [Document(page_content=text, metadata={"source": source, "title": str(title)})] if text else []


def _load_docx(path: Path, source: str) -> list[Document]:
    import docx  # python-docx

    document = docx.Document(str(path))
    text = clean_text("\n".join(p.text for p in document.paragraphs))
    return [Document(page_content=text, metadata={"source": source, "title": _title_from_path(path)})] if text else []


def _load_plain(path: Path, source: str) -> list[Document]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    return [Document(page_content=text, metadata={"source": source, "title": _title_from_path(path)})] if text else []


def load_file(path: Path, root: Path) -> list[Document]:
    source = path.relative_to(root).as_posix()
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path, source)
    if suffix in {".html", ".htm"}:
        return _load_html(path, source)
    if suffix == ".docx":
        return _load_docx(path, source)
    return _load_plain(path, source)


def _load_safely(args: tuple[str, str]) -> tuple[str, list[Document], str | None]:
    """Worker entry point: top-level so it pickles for ProcessPoolExecutor."""
    path, root = Path(args[0]), Path(args[1])
    try:
        return path.relative_to(root).as_posix(), load_file(path, root), None
    except Exception as exc:
        return path.relative_to(root).as_posix(), [], f"{exc.__class__.__name__}: {exc}"


def _resolve_workers(workers: int | None, paths: list[Path]) -> int:
    requested = config.INGEST_WORKERS if workers is None else workers
    if requested > 0:
        return min(requested, len(paths))
    total = sum(p.stat().st_size for p in paths)
    if total < _PARALLEL_MIN_BYTES:
        return 1
    return max(1, min(os.cpu_count() or 1, 8, len(paths)))


def load_directory(data_dir: str | Path, workers: int | None = None) -> list[Document]:
    root = Path(data_dir).resolve()
    if not root.is_dir():
        raise SystemExit(f"Data directory not found: {root}")

    paths = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES)
    if not paths:
        raise SystemExit(
            f"No supported documents in {root}. "
            f"Drop some {', '.join(sorted(SUPPORTED_SUFFIXES))} files in there and re-run."
        )

    n_workers = _resolve_workers(workers, paths)
    jobs = [(str(p), str(root)) for p in paths]
    if n_workers > 1:
        # map() preserves input order, so the corpus is identical to a serial run.
        with ProcessPoolExecutor(max_workers=n_workers) as pool:
            results = list(pool.map(_load_safely, jobs, chunksize=1))
    else:
        results = [_load_safely(job) for job in jobs]

    docs: list[Document] = []
    for rel, loaded, error in results:
        if error:
            logger.warning("Skipping %s: %s", rel, error)
            continue
        logger.debug("Loaded %-45s %3d section(s)", rel, len(loaded))
        docs.extend(loaded)

    logger.info("Loaded %d section(s) from %d file(s) with %d worker(s)", len(docs), len(paths), n_workers)
    return docs


def corpus_stats(docs: list[Document]) -> dict:
    """What was actually ingested - the numbers the README quotes."""
    return {
        "files": len({d.metadata.get("source") for d in docs}),
        "pdf_pages": sum(1 for d in docs if d.metadata.get("page")),
        "sections": len(docs),
        "chars": sum(len(d.page_content) for d in docs),
    }
